#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


NS = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
BLOCKING_FILL = PatternFill("solid", fgColor="FFFFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFFFF2CC")

DETECTION_REUSE_DIMS = {"software", "vendor", "attack_name", "attack_type"}
CODE_DIMS = {
    "cve",
    "cwe_nocn",
    "capec_nocn",
    "mitre_techniques",
    "mitre_tactics",
    "mitre_mitigation",
    "ics_mitre_techniques",
    "ics_mitre_tactics",
    "nist_control",
    "nist_control_nocn",
}


def clean(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\u00a0", " ").strip()


def norm(value: object) -> str:
    text = clean(value).lower()
    text = text.replace("：", ":").replace("（", "(").replace("）", ")")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def compact(value: object) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", norm(value))


def code(value: str, pattern: str) -> str:
    m = re.search(pattern, clean(value), flags=re.I)
    return m.group(0).upper() if m else ""


def is_ics_tactic(value: str) -> bool:
    c = code(value, r"TA\d{4}")
    return bool(c and "TA0100" <= c <= "TA0111")


def is_ics_technique(value: str) -> bool:
    c = code(value, r"T\d{4}(?:\.\d{3})?")
    base = c.split(".", 1)[0] if c else ""
    return bool(base and "T0800" <= base <= "T0895")


def colnum(ref: str) -> int:
    m = re.match(r"([A-Z]+)", ref or "")
    if not m:
        return 0
    n = 0
    for ch in m.group(1):
        n = n * 26 + ord(ch) - 64
    return n


def read_xlsx_first_sheet_xml(path: Path) -> list[dict[str, str]]:
    with ZipFile(path) as zf:
        workbook = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        rel_map = {
            rel.attrib["Id"]: rel.attrib["Target"].lstrip("/")
            for rel in rels
            if rel.attrib.get("Type", "").endswith("/worksheet")
        }
        first_sheet = workbook.find(".//a:sheet", NS)
        if first_sheet is None:
            return []
        rid = first_sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        sheet_path = rel_map[rid]
        if not sheet_path.startswith("xl/"):
            sheet_path = f"xl/{sheet_path}"

        shared: list[str] = []
        if "xl/sharedStrings.xml" in zf.namelist():
            root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in root.findall("a:si", NS):
                shared.append("".join(t.text or "" for t in si.findall(".//a:t", NS)))

        ws = ET.fromstring(zf.read(sheet_path))
        rows: list[list[str]] = []
        for row in ws.findall(".//a:sheetData/a:row", NS):
            values: dict[int, str] = {}
            for cell in row.findall("a:c", NS):
                idx = colnum(cell.attrib.get("r", ""))
                if idx <= 0:
                    continue
                v = cell.find("a:v", NS)
                if v is None:
                    value = ""
                elif cell.attrib.get("t") == "s":
                    value = shared[int(v.text)]
                else:
                    value = v.text or ""
                values[idx] = clean(value)
            if values:
                rows.append([values.get(i, "") for i in range(1, max(values) + 1)])
    if not rows:
        return []
    headers = [clean(x) for x in rows[0]]
    records = []
    for row in rows[1:]:
        rec = {headers[i]: clean(row[i]) if i < len(row) else "" for i in range(len(headers))}
        if any(rec.values()):
            records.append(rec)
    return records


def read_source_tags(path: Path) -> list[dict[str, str]]:
    # XML parsing is intentional: some exports have stale worksheet dimension metadata.
    return read_xlsx_first_sheet_xml(path)


def read_language(path: Path | None) -> dict[str, dict[str, str]]:
    if not path:
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    out: dict[str, dict[str, str]] = {}
    for ws in wb.worksheets:
        headers = [clean(c.value) for c in ws[1]]
        if "uuid" not in headers:
            continue
        idx = {h: i for i, h in enumerate(headers)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            uuid = clean(row[idx["uuid"]]).lower()
            if not uuid:
                continue
            item = out.setdefault(uuid, {})
            for target, aliases in {
                "rule_name": ["cn_name", "name", "rule_name_cn"],
                "rule_name_en": ["en_name", "name_en", "rule_name_en"],
            }.items():
                for alias in aliases:
                    if alias in idx:
                        value = clean(row[idx[alias]])
                        if value:
                            item[target] = value
                            break
    return out


def strip_prefix(value: str) -> str:
    text = clean(value)
    for prefix in ["ATT&CK:", "NIST:", "CWE:", "CAPEC:", "Control:", "OS:", "RunAs:"]:
        if text.lower().startswith(prefix.lower()):
            return text[len(prefix):].strip()
    return text


def infer_source_dim(raw_type: str, raw_value: str) -> tuple[str, str] | None:
    tag_type = clean(raw_type).lower()
    value = clean(raw_value)
    lower = value.lower()
    if not value:
        return None

    if tag_type:
        value = strip_prefix(value)
        if tag_type in {"mitre_tactics", "official_mitre_tactics"} and is_ics_tactic(value):
            return "ics_mitre_tactics", code(value, r"TA\d{4}")
        if tag_type in {"mitre_techniques", "official_mitre_techniques"} and is_ics_technique(value):
            return "ics_mitre_techniques", code(value, r"T\d{4}(?:\.\d{3})?")
        if tag_type == "cve":
            c = code(value, r"CVE-\d{4}-\d+")
            return ("cve", c) if c else None
        if tag_type == "cwe_nocn":
            n = re.search(r"\d+", value)
            return ("cwe_nocn", f"CWE-{n.group(0)}") if n else None
        return tag_type, value

    structured = re.match(
        r"^(control|os|run_as|src_destination|mitre_tactics|mitre_techniques|ics_mitre_tactics|ics_mitre_techniques|mitre_mitigation|nist_control|cve|cwe_nocn|capec_nocn|campaign_nocn)\s+(.+)$",
        value,
        flags=re.I,
    )
    if structured:
        dim = structured.group(1).lower()
        body = structured.group(2).strip()
        if dim == "mitre_tactics":
            tac = code(body, r"TA\d{4}") or body
            return ("ics_mitre_tactics", tac) if is_ics_tactic(tac) else (dim, tac)
        if dim == "mitre_techniques":
            tech = code(body, r"T\d{4}(?:\.\d{3})?") or body
            return ("ics_mitre_techniques", tech) if is_ics_technique(tech) else (dim, tech)
        if dim == "ics_mitre_tactics":
            return dim, code(body, r"TA\d{4}") or body
        if dim == "ics_mitre_techniques":
            return dim, code(body, r"T\d{4}(?:\.\d{3})?") or body
        if dim == "mitre_mitigation":
            return dim, code(body, r"M\d{4}") or body
        if dim == "cwe_nocn":
            n = re.search(r"\d+", body)
            return dim, f"CWE-{n.group(0)}" if n else body
        if dim == "capec_nocn":
            n = re.search(r"\d+", body)
            return dim, f"CAPEC-{n.group(0)}" if n else body
        if dim == "cve":
            c = code(body, r"CVE-\d{4}-\d+")
            return (dim, c) if c else None
        return dim, strip_prefix(body)

    if lower.startswith("att&ck:"):
        body = strip_prefix(value)
        if re.fullmatch(r"TA\d{4}", body, flags=re.I):
            tac = body.upper()
            return ("ics_mitre_tactics", tac) if is_ics_tactic(tac) else ("mitre_tactics", tac)
        if re.fullmatch(r"T\d{4}(?:\.\d{3})?", body, flags=re.I):
            tech = body.upper()
            return ("ics_mitre_techniques", tech) if is_ics_technique(tech) else ("mitre_techniques", tech)
        if re.fullmatch(r"M\d{4}", body, flags=re.I):
            return "mitre_mitigation", body.upper()
    if lower.startswith("control:"):
        return "control", strip_prefix(value)
    if lower.startswith("os:"):
        return "os", strip_prefix(value)
    if lower.startswith("runas:"):
        return "run_as", strip_prefix(value)
    if lower.startswith("src:") or lower.startswith("dst:"):
        return "src_destination", value
    if lower.startswith("nist:"):
        return "nist_control", strip_prefix(value)
    if lower.startswith("cwe:") or re.fullmatch(r"CWE[-:]\d+", value, flags=re.I):
        n = re.search(r"\d+", value)
        return ("cwe_nocn", f"CWE-{n.group(0)}") if n else None
    if lower.startswith("capec:") or re.fullmatch(r"CAPEC[-:]\d+", value, flags=re.I):
        n = re.search(r"\d+", value)
        return ("capec_nocn", f"CAPEC-{n.group(0)}") if n else None
    if re.fullmatch(r"CVE-\d{4}-\d+", value, flags=re.I):
        return "cve", value.upper()
    return None


def read_output(path: Path) -> tuple[dict[str, dict[str, str]], dict[str, dict[str, set[str]]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if "summary" not in wb.sheetnames:
        raise SystemExit(f"output workbook lacks summary sheet: {path}")

    summary: dict[str, dict[str, str]] = {}
    ws = wb["summary"]
    headers = [clean(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    if "uuid" not in idx:
        raise SystemExit("summary sheet lacks uuid")
    for row in ws.iter_rows(min_row=2, values_only=True):
        uuid = clean(row[idx["uuid"]]).lower()
        if not uuid:
            continue
        summary[uuid] = {h: clean(row[i]) if i < len(row) else "" for h, i in idx.items()}

    tags: dict[str, dict[str, set[str]]] = defaultdict(lambda: defaultdict(set))
    for ws in wb.worksheets:
        if ws.title == "summary":
            continue
        headers = [clean(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = {h: i for i, h in enumerate(headers)}
        if not {"uuid", "tag_cn"}.issubset(idx):
            continue
        dim = ws.title
        if "tag_type" in idx:
            dim_values = set()
        else:
            dim_values = {dim}
        for row in ws.iter_rows(min_row=2, values_only=True):
            uuid = clean(row[idx["uuid"]]).lower()
            if not uuid:
                continue
            row_dim = clean(row[idx["tag_type"]]).lower() if "tag_type" in idx else dim
            cn = clean(row[idx["tag_cn"]])
            en = clean(row[idx["tag_en"]]) if "tag_en" in idx else ""
            if cn:
                tags[uuid][row_dim].add(cn)
            if en:
                tags[uuid][row_dim].add(en)
    return summary, tags


def read_detection_master(path: Path | None, target_uuids: set[str]) -> dict[str, dict[str, set[str]]]:
    if not path:
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    out: dict[str, dict[str, set[str]]] = defaultdict(lambda: defaultdict(set))
    for dim in DETECTION_REUSE_DIMS:
        if dim not in wb.sheetnames:
            continue
        ws = wb[dim]
        headers = [clean(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = {h: i for i, h in enumerate(headers)}
        if "uuid" not in idx:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            uuid = clean(row[idx["uuid"]]).lower()
            if uuid not in target_uuids:
                continue
            for col in ["tag_cn", "tag_en", "tag"]:
                if col in idx:
                    value = clean(row[idx[col]])
                    if value:
                        out[uuid][dim].add(value)
    return out


def value_present(expected: str, actual_values: Iterable[str], dim: str) -> bool:
    exp = clean(expected)
    if not exp:
        return True
    actual = [clean(v) for v in actual_values if clean(v)]
    if dim in CODE_DIMS or re.fullmatch(r"(?:CVE-\d{4}-\d+|CWE-\d+|CAPEC-\d+|TA\d{4}|T\d{4}(?:\.\d{3})?|M\d{4}|[A-Z]{2}-\d{2})", exp, flags=re.I):
        exp_norm = normalize_code(exp)
        actual_codes = set()
        for value in actual:
            actual_codes.add(normalize_code(value))
            actual_codes.update(extract_codes(value))
        return exp_norm in actual_codes
    cexp = compact(exp)
    return any(cexp == compact(a) or cexp in compact(a) or compact(a) in cexp for a in actual)


def normalize_code(value: str) -> str:
    text = clean(value).upper().replace("：", ":")
    text = re.sub(r"\bCWE[:：](\d+)\b", r"CWE-\1", text)
    text = re.sub(r"\bCAPEC[:：](\d+)\b", r"CAPEC-\1", text)
    for pattern in [
        r"CVE-\d{4}-\d+",
        r"CWE-\d+",
        r"CAPEC-\d+",
        r"TA\d{4}",
        r"T\d{4}(?:\.\d{3})?",
        r"M\d{4}",
        r"[A-Z]{2}-\d{2}",
    ]:
        found = code(text, pattern)
        if found:
            return found
    return text


def extract_codes(value: str) -> set[str]:
    text = clean(value).upper().replace("：", ":")
    text = re.sub(r"\bCWE[:：](\d+)\b", r"CWE-\1", text)
    text = re.sub(r"\bCAPEC[:：](\d+)\b", r"CAPEC-\1", text)
    patterns = [
        r"CVE-\d{4}-\d+",
        r"CWE-\d+",
        r"CAPEC-\d+",
        r"TA\d{4}",
        r"T\d{4}(?:\.\d{3})?",
        r"M\d{4}",
        r"[A-Z]{2}-\d{2}",
    ]
    out = set()
    for pattern in patterns:
        out.update(m.group(0).upper() for m in re.finditer(pattern, text))
    return out


def add_issue(issues: list[dict[str, str]], severity: str, issue_type: str, uuid: str, detail: str, expected: str = "", actual: str = "") -> None:
    issues.append(
        {
            "severity": severity,
            "issue_type": issue_type,
            "uuid": uuid,
            "detail": detail,
            "expected": expected,
            "actual": actual,
        }
    )


def write_report(path: Path, summary: dict[str, int | str], issues: list[dict[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(["metric", "value"])
    for k, v in summary.items():
        ws.append([k, v])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 28

    iws = wb.create_sheet("issues")
    headers = ["severity", "issue_type", "uuid", "detail", "expected", "actual"]
    iws.append(headers)
    for cell in iws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for issue in issues:
        iws.append([issue[h] for h in headers])
        fill = BLOCKING_FILL if issue["severity"] == "BLOCKING" else WARN_FILL
        for cell in iws[iws.max_row]:
            cell.fill = fill
    widths = [14, 28, 38, 80, 55, 55]
    for idx, width in enumerate(widths, 1):
        iws.column_dimensions[get_column_letter(idx)].width = width
    iws.auto_filter.ref = iws.dimensions
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Reconcile Validation tag output with source tags, language workbook, and optional Detection master.")
    parser.add_argument("--source-tags", required=True, type=Path)
    parser.add_argument("--language", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--detection-master", type=Path)
    parser.add_argument("--report", required=True, type=Path)
    args = parser.parse_args()

    source_records = read_source_tags(args.source_tags)
    language = read_language(args.language)
    output_summary, output_tags = read_output(args.output)

    source_uuids = {clean(r.get("uuid")).lower() for r in source_records if clean(r.get("uuid"))}
    output_uuids = set(output_summary)
    language_uuids = set(language)
    issues: list[dict[str, str]] = []

    for uuid in sorted(source_uuids - output_uuids):
        add_issue(issues, "BLOCKING", "missing_output_uuid", uuid, "Source UUID is absent from output")
    for uuid in sorted(output_uuids - source_uuids):
        add_issue(issues, "WARN", "extra_output_uuid", uuid, "Output UUID is absent from source tag workbook")

    for uuid in sorted(output_uuids & language_uuids):
        row = output_summary[uuid]
        lang = language[uuid]
        for field in ["rule_name", "rule_name_en"]:
            expected = clean(lang.get(field))
            actual = clean(row.get(field))
            if expected and actual and norm(expected) != norm(actual):
                add_issue(issues, "BLOCKING", f"language_{field}_mismatch", uuid, "Output metadata differs from language workbook", expected, actual)
            elif expected and not actual:
                add_issue(issues, "BLOCKING", f"language_{field}_missing", uuid, "Output metadata missing language workbook value", expected, actual)

    expected_by_uuid_dim: dict[str, dict[str, set[str]]] = defaultdict(lambda: defaultdict(set))
    for rec in source_records:
        uuid = clean(rec.get("uuid")).lower()
        inferred = infer_source_dim(rec.get("tag_type", ""), rec.get("tag", ""))
        if uuid and inferred:
            dim, value = inferred
            expected_by_uuid_dim[uuid][dim].add(value)

    for uuid, dims in expected_by_uuid_dim.items():
        for dim, values in dims.items():
            actual_values = output_tags.get(uuid, {}).get(dim, set())
            # nist_control may be exported as nist_control_nocn in older files.
            if dim == "nist_control" and not actual_values:
                actual_values = output_tags.get(uuid, {}).get("nist_control_nocn", set())
            for expected in values:
                if not value_present(expected, actual_values, dim):
                    add_issue(issues, "BLOCKING", "source_structured_tag_missing", uuid, f"Expected source structured tag missing from output dimension {dim}", expected, " | ".join(sorted(actual_values)))

    cve_pat = re.compile(r"CVE-\d{4}-\d+", re.I)
    for uuid, row in output_summary.items():
        title_text = " ".join([row.get("rule_name", ""), row.get("rule_name_en", "")])
        expected_cves = {m.group(0).upper() for m in cve_pat.finditer(title_text)}
        actual_cves = {v.upper() for v in output_tags.get(uuid, {}).get("cve", set()) if cve_pat.fullmatch(v)}
        for cve in sorted(expected_cves - actual_cves):
            add_issue(issues, "BLOCKING", "title_cve_missing", uuid, "CVE in output title is missing from cve tags", cve, " | ".join(sorted(actual_cves)))
        for cve in sorted(actual_cves - expected_cves):
            # Source workbook may contain CVE rows even when title is incomplete; warn rather than block.
            if cve not in expected_by_uuid_dim.get(uuid, {}).get("cve", set()):
                add_issue(issues, "WARN", "output_cve_not_in_title_or_source", uuid, "Output CVE is not found in title or source structured rows", "", cve)

    detection = read_detection_master(args.detection_master, output_uuids) if args.detection_master else {}
    for uuid, dims in detection.items():
        for dim, values in dims.items():
            actual_values = output_tags.get(uuid, {}).get(dim, set())
            for expected in values:
                if not value_present(expected, actual_values, dim):
                    add_issue(issues, "BLOCKING", "detection_reuse_missing", uuid, f"Detection master value missing from Validation output dimension {dim}", expected, " | ".join(sorted(actual_values)))

    severity_counts = Counter(issue["severity"] for issue in issues)
    summary = {
        "source_rows": len(source_records),
        "source_uuids": len(source_uuids),
        "language_uuids": len(language_uuids),
        "output_uuids": len(output_uuids),
        "blocking_issues": severity_counts.get("BLOCKING", 0),
        "warn_issues": severity_counts.get("WARN", 0),
        "report": str(args.report),
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    write_report(args.report, summary, issues)
    print("validation_reconciliation")
    for k, v in summary.items():
        print(f"{k}={v}")
    return 2 if severity_counts.get("BLOCKING", 0) else 0


if __name__ == "__main__":
    sys.exit(main())
